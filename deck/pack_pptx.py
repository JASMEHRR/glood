"""Packs the rendered deck into a 16:9 .pptx: a background picture per slide (the artwork, text hidden)
with every piece of text laid on top as an editable text box, plus speaker notes.

Positions arrive in CSS pixels of a 1600 x 900 slide.
"""
import json
import math
import re
import sys

from lxml import etree
from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SLIDE_W, SLIDE_H = 9144000, 5143500  # 10 x 5.625 in, same as the original deck
PX = SLIDE_W / 1600  # EMU per CSS pixel
PT = 720 / 1600  # points per CSS pixel
ALIGN = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT}


def rgb(css):
    nums = [float(v) for v in re.findall(r"[\d.]+", css)]
    return RGBColor(*[int(v) for v in nums[:3]]), (nums[3] if len(nums) > 3 else 1.0)


def split_rupee(run):
    """Lilita One has no rupee sign: set that one character in Space Grotesk Bold."""
    if run["font"] != "Lilita One" or "₹" not in run["text"]:
        return [run]
    out = []
    for part in re.split("(₹)", run["text"]):
        if part:
            out.append({**run, "text": part, **({"font": "Space Grotesk", "weight": 700} if part == "₹" else {})})
    return out


def add_outline(rpr, width_px):
    ln = etree.SubElement(rpr, qn("a:ln"), w=str(int(width_px * PX)))
    fill = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr"), val="000000")
    # a:ln must come before the fill elements in rPr
    rpr.remove(ln)
    rpr.insert(0, ln)


def add_shadow(rpr, css):
    m = re.search(r"(-?[\d.]+)px (-?[\d.]+)px", css)
    if not m:
        return
    dx, dy = float(m.group(1)), float(m.group(2))
    effects = etree.SubElement(rpr, qn("a:effectLst"))
    latin = rpr.find(qn("a:latin"))  # schema order: ln, fill, effectLst, ..., latin
    if latin is not None:
        latin.addprevious(effects)
    sh = etree.SubElement(effects, qn("a:outerShdw"), blurRad="0", dist=str(int(math.hypot(dx, dy) * PX)),
                          dir=str(int(math.degrees(math.atan2(dy, dx)) * 60000) % 21600000), algn="tl", rotWithShape="0")
    etree.SubElement(sh, qn("a:srgbClr"), val="000000")


def add_text(slide, t):
    pl, pt, pr, pb = t["pad"]
    bl, bt, br, bb = t["border"]
    single = t["h"] - pt - pb - bt - bb < t["lineHeight"] * 1.6
    w = t["w"] - bl - br + (0 if single else t["w"] * 0.03)  # a little slack so wrapping matches the browser
    h = t["h"] - bt - bb
    box = slide.shapes.add_textbox(Emu(int((t["cx"] - w / 2) * PX)), Emu(int((t["cy"] - h / 2) * PX)),
                                   Emu(int(w * PX)), Emu(int(h * PX)))
    box.rotation = t["rot"]
    tf = box.text_frame
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Emu(int(v * PX)) for v in (pl, pt, pr, pb))
    tf.word_wrap = not single
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if single else MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = ALIGN.get(t["align"], PP_ALIGN.LEFT)
    p.line_spacing = max(0.8, min(1.6, t["lineHeight"] / t["fontSize"] / 1.2))
    for raw in t["runs"]:
        for run in split_rupee(raw):
            r = p.add_run()
            text = run["text"]
            if run["font"] == "Bungee":
                text = text.upper()  # Bungee is a capitals-only face
            r.text = text
            f = r.font
            f.name = run["font"]
            f.size = Pt(round(run["size"] * PT, 1))
            f.bold = run["weight"] >= 600
            color, alpha = rgb(run["color"])
            f.color.rgb = color
            rpr = r._r.get_or_add_rPr()
            if run["stroke"]:
                add_outline(rpr, run["stroke"])
            if run["shadow"]:
                add_shadow(rpr, run["shadow"])


TINT_ALPHA = 72  # % opacity of the colour panel over the backdrop on content slides

# Entrance effects (PowerPoint MsoAnimEffect ids) and durations in seconds, per kind of item.
EFFECT = {
    "sticker": (26, 0.6),  # Bounce: stickers slap on, like the laptop story on the site
    "card": (23, 0.45),  # Zoom
    "chip": (31, 0.4),  # Grow & Turn
    "text": (42, 0.5),  # Float In, for headings and other free text
    "inner": (10, 0.3),  # Fade, for text that arrives with its card or chip
}


def add_tint(slide, css):
    color, _ = rgb(css)
    rect = slide.shapes.add_shape(1, 0, 0, Emu(SLIDE_W), Emu(SLIDE_H))  # 1 = rectangle
    rect.name = "!!Tint"  # the !! prefix makes Morph match it across slides, so the colour glides
    rect.line.fill.background()
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    clr = rect.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    etree.SubElement(clr, qn("a:alpha"), val=str(TINT_ALPHA * 1000))
    rect.shadow.inherit = False


MORPH = (
    '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
    '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
    '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow">'
    '<p159:morph option="byObject"/></p:transition></mc:Choice>'
    '<mc:Fallback><p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow">'
    "<p:fade/></p:transition></mc:Fallback></mc:AlternateContent>"
)


def add_morph(slide):
    """Morph transition (falls back to a fade on old PowerPoint). Goes right after clrMapOvr."""
    slide._element.find(qn("p:clrMapOvr")).addnext(etree.fromstring(MORPH))


def picture(slide, image, x, y, w, h, name):
    pic = slide.shapes.add_picture(image, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    pic.name = name
    return pic


def plan_for(s):
    """Entrance order: everything in reading order, one after another; text inside a card or chip
    fades in together with it, and a run of chips arrives together with a small stagger."""
    items = [("pop", i, p["order"]) for i, p in enumerate(s["pops"])]
    items += [("text", j, t["order"]) for j, t in enumerate(s["texts"]) if t["pop"] < 0]
    steps, prev_chip, stagger = [], False, 0
    for kind, i, _ in sorted(items, key=lambda it: it[2]):
        if kind == "pop":
            pk = s["pops"][i]["kind"]
            effect, dur = EFFECT[pk]
            chip = pk == "chip"
            stagger = stagger + 0.08 if chip and prev_chip else 0
            steps.append({"shape": f"Pop {i}", "effect": effect, "duration": dur,
                          "trigger": "with" if chip and prev_chip else "after", "delay": stagger})
            for j, t in enumerate(s["texts"]):
                if t["pop"] == i:
                    e, d = EFFECT["inner"]
                    steps.append({"shape": f"Text {j}", "effect": e, "duration": d, "trigger": "with", "delay": stagger + 0.15})
            prev_chip = chip
        else:
            e, d = EFFECT["text"]
            steps.append({"shape": f"Text {i}", "effect": e, "duration": d, "trigger": "after", "delay": 0})
            prev_chip = False
    return steps


data, out, plan_out = json.load(open(sys.argv[1], encoding="utf-8")), sys.argv[2], sys.argv[3]
# Content slides sit on a soft-focus copy of the landing page, so text stays clean while the logo
# colours still glow through; Morph blurs it in when leaving the sharp cover.
soft = data["backdrop"].replace(".png", "-soft.png")
Image.open(data["backdrop"]).filter(ImageFilter.GaussianBlur(18)).save(soft)
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
blank = prs.slide_layouts[6]
count, plans = 0, []
for s in data["slides"]:
    slide = prs.slides.add_slide(blank)
    picture(slide, data["backdrop"] if s["cover"] else soft, 0, 0, 1600, 900, "!!Backdrop")
    if not s["cover"]:
        add_tint(slide, s["tint"])
    picture(slide, s["art"], 0, 0, 1600, 900, "Artwork")
    for i, p in enumerate(s["pops"]):
        picture(slide, p["image"], p["x"], p["y"], p["w"], p["h"], f"Pop {i}")
    for j, t in enumerate(s["texts"]):
        add_text(slide, t)
        slide.shapes[-1].name = f"Text {j}"
        count += 1
    slide.notes_slide.notes_text_frame.text = s["notes"] or ""
    add_morph(slide)
    plans.append(plan_for(s))
prs.save(out)
json.dump(plans, open(plan_out, "w", encoding="utf-8"), indent=1)
print(f"Saved {len(plans)} slides, {count} editable text boxes, {sum(len(p) for p in plans)} animations planned")
